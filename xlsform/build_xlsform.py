"""
build_xlsform.py

Generates criminal_risk_assessment.xlsx (an ODK XLSForm) that digitizes the
Manitoba Families "Criminal Risk Assessment Request" paper form.

Run:
    python3 build_xlsform.py

This script is the source of truth for the form -- re-run it any time the
field list below changes, instead of hand-editing the .xlsx.
"""

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

SURVEY_COLUMNS = [
    "type", "name", "label", "required", "required_message", "appearance",
    "hint", "relevant", "default", "constraint", "constraint_message",
    "calculation", "repeat_count", "read_only",
]

CHOICES_COLUMNS = ["list name", "name", "label"]
SETTINGS_COLUMNS = ["form_title", "form_id", "version", "default_language"]

CONSENT_TEXT = (
    "As a person who has, or may have, care, custody, control or charge of a "
    "child in receipt of services under The Child and Family Services Act, I "
    "authorize the Criminal Risk Assessment Unit of Manitoba Families' Child "
    "Protection Branch (\"CRAU\") to conduct enquiries of the Winnipeg Police "
    "Service (WPS), the RCMP and other law enforcement agencies necessary to "
    "assess the risk that I may endanger the life, health or emotional "
    "wellbeing of a child. I understand that this information is requested "
    "by (CFS Agency) for the purposes of, and in accordance with, s. 18.4(1.1) "
    "of The Child and Family Services Act, and may include a criminal record, "
    "criminal and Provincial Act convictions, orders or charges, other "
    "involvement/contact with law enforcement (including non-conviction "
    "information) or other information. I authorize the disclosure of said "
    "information to CRAU and an authorized Child and Family Services Agency "
    "designate or designates. I also authorize the disclosure of the personal "
    "identifying information set out below to CRAU, the WPS, RCMP and other "
    "law enforcement agencies for the purpose of completing a Criminal Risk "
    "Assessment. I understand that the results of this Criminal Risk "
    "Assessment are confidential, and may not be provided to me, but may be "
    "disclosed in accordance with s.76 of the Child and Family Services Act."
)

PAGE2_LEGAL_TEXT = (
    "The identifying information supplied on page 1 will be researched "
    "through the Canadian Police Information Centre (CPIC) and Winnipeg "
    "Police Service Records Management System (NICHE). The research will NOT "
    "include a vulnerable sector search or information in respect of "
    "pardons, which information may impact the results of a risk assessment. "
    "Criminal Convictions are strictly confidential and cannot be shared "
    "without written authorization of the person involved. Personal "
    "Information is governed under Section 8 of the Privacy Act (Federal "
    "Statute). Record information is based on NAME AND DATE OF BIRTH SEARCH "
    "ONLY and does not necessarily indicate subject involvement. Verification "
    "can only be provided through the submission of fingerprints to the "
    "local law enforcement agency for the area in which the person being "
    "assessed resides."
)

FINAL_NOTE_TEXT = (
    "NOTE: The assessment completed by the Criminal Risk Assessment Unit of "
    "the Department of Families Child Protection Branch does not replace a "
    "criminal records check."
)

# Each row is a dict keyed by SURVEY_COLUMNS; missing keys default to "".
survey_rows = [
    # ---------------------------------------------------------------
    # PAGE 1 - Consent block
    # ---------------------------------------------------------------
    dict(type="begin group", name="page1_consent",
         label="Consent for Criminal Risk Assessment and Release of Information",
         appearance="field-list"),

    dict(type="note", name="consent_notice", label=CONSENT_TEXT),

    dict(type="select_one consent_status", name="consent_status",
         label="Consent Status", required="yes",
         hint="The paper form has a single 'Unconsented' checkbox; the "
              "implied opposite state (consent given) is captured here as "
              "'Consented' so the field is answerable either way."),

    dict(type="date", name="consent_date", label="Date", required="yes"),

    dict(type="image", name="signature_person_assessed",
         label="Signature of Person Being Assessed",
         appearance="signature max-pixels:640",
         relevant="${consent_status}='consented'", required="yes",
         hint="Captured as a signature image, same pattern used for the "
              "witness/patient signature in the reference sample form."),

    dict(type="text", name="witness_name",
         label="Witness (if consenting)",
         relevant="${consent_status}='consented'"),

    dict(type="end group"),

    # ---------------------------------------------------------------
    # PAGE 1 - Identifying information (fields 1-10)
    # ---------------------------------------------------------------
    dict(type="begin group", name="page1_identity",
         label="Identifying Information (please print clearly)",
         appearance="field-list"),

    dict(type="text", name="first_name", label="1. First Name", required="yes"),
    dict(type="text", name="second_name", label="2. Second Name"),
    dict(type="text", name="last_name", label="3. Last Name", required="yes"),
    dict(type="date", name="date_of_birth", label="4. Date of Birth", required="yes"),

    dict(type="select_one gender", name="gender", label="5. Sex", required="yes",
         hint="Source PDF's checkbox row reads 'MALE [] MALE []', which "
              "looks like an OCR/typesetting duplication of the source "
              "document rather than an intentional field. Modelled here as "
              "the evident intent: Male / Female. Flagged as an assumption "
              "in the README."),

    dict(type="text", name="other_last_names", label="6. Other Last Names Used"),
    dict(type="text", name="other_first_names",
         label="7. Other First Names Used / Also Goes By"),
    dict(type="text", name="current_address",
         label="8. Current Address (include postal code)"),
    dict(type="text", name="current_phone", label="9. Current Phone Number(s)"),
    dict(type="text", name="birth_place",
         label="10. City/Province or Country of Birth"),

    dict(type="calculate", name="full_name_calc",
         calculation="concat(${first_name}, ' ', ${last_name})"),

    dict(type="end group"),

    # ---------------------------------------------------------------
    # PAGE 1 - Identity verification (2 pieces of ID)
    # ---------------------------------------------------------------
    dict(type="begin group", name="page1_id_verification",
         label="Identity Verification", appearance="field-list"),

    dict(type="note", name="id_verification_note",
         label="Subject's name must be identified with TWO PIECES OF "
               "IDENTIFICATION (MB driver's licence & photo ID is "
               "preferable)."),

    dict(type="select_multiple id_type", name="id_types_provided",
         label="Identification Provided", required="yes",
         constraint="count-selected(.) >= 2",
         constraint_message="Please select at least TWO pieces of "
                             "identification, as required on the form."),

    dict(type="text", name="other_id_specify", label="Other (specify ID)",
         relevant="selected(${id_types_provided}, 'other')", required="yes"),

    dict(type="text", name="mb_dl_licence_number",
         label="MB Driver's Licence Number (section 4d on licence)",
         relevant="selected(${id_types_provided}, 'mb_dl')", required="yes"),

    dict(type="end group"),

    # ---------------------------------------------------------------
    # PAGE 2
    # ---------------------------------------------------------------
    dict(type="begin group", name="page2_request",
         label="Criminal Risk Assessment Unit - Request Details (Page 2)",
         appearance="field-list"),

    dict(type="note", name="page2_legal_note", label=PAGE2_LEGAL_TEXT),

    dict(type="text", name="assessed_person_name_display",
         label="Name of Person Being Assessed",
         hint="*Must match information on page 1 - auto-filled from the "
              "First/Last Name captured above; not re-typed by the user.",
         calculation="${full_name_calc}", read_only="true"),

    dict(type="text", name="agency_name",
         label="*Name of Agency Submitting Request", required="yes"),

    dict(type="select_one reason", name="reason_for_assessment",
         label="*Reason for Risk Assessment", required="yes"),

    dict(type="note", name="consent_required_warning",
         label="This reason requires consent, but page 1 records the "
               "person as 'Unconsented'. Please confirm this is correct "
               "before submitting.",
         relevant="(${reason_for_assessment}='place_of_safety' or "
                  "${reason_for_assessment}='kinship_customary_care') and "
                  "${consent_status}='unconsented'"),

    dict(type="text", name="assigned_worker", label="*Assigned Worker",
         required="yes"),

    dict(type="date", name="last_assessment_date",
         label="Date of Last Criminal Risk Assessment (if known)"),

    dict(type="text", name="submitting_designate",
         label="*Submitting Designate", required="yes"),

    dict(type="text", name="designate_phone", label="*Designate Phone Number",
         required="yes"),

    dict(type="text", name="designate_email", label="*Designate Email",
         required="yes",
         constraint="regex(., '^[^@\\s]+@[^@\\s]+\\.[^@\\s]+$')",
         constraint_message="Enter a valid email address."),

    dict(type="text", name="designate_fax", label="Designate Fax Number"),

    dict(type="date", name="request_date", label="*Request Date",
         required="yes", default="today()"),

    dict(type="note", name="final_note", label=FINAL_NOTE_TEXT),

    dict(type="end group"),
]

choices_rows = [
    ("consent_status", "consented", "Consented"),
    ("consent_status", "unconsented", "Unconsented"),

    ("gender", "male", "Male"),
    ("gender", "female", "Female"),

    ("id_type", "birth_certificate", "Birth Certificate"),
    ("id_type", "sin_card", "Social Insurance Card"),
    ("id_type", "mb_health_card", "Manitoba Health Card"),
    ("id_type", "treaty_card", "Treaty Card"),
    ("id_type", "mb_dl", "MB Driver's Licence with Photo"),
    ("id_type", "other", "Other (specify)"),

    ("reason", "child_protection_concerns",
     "Child Protection Concerns (with or without consent)"),
    ("reason", "place_of_safety", "Place of Safety (consent required)"),
    ("reason", "kinship_customary_care",
     "Kinship or Customary Care Agreement (consent required)"),
]

settings_row = {
    "form_title": "Criminal Risk Assessment Request",
    "form_id": "criminal_risk_assessment_request",
    "version": "2025011001",
    "default_language": "English (en)",
}


def build():
    wb = openpyxl.Workbook()

    # --- survey sheet ---
    ws = wb.active
    ws.title = "survey"
    ws.append(SURVEY_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in survey_rows:
        ws.append([row.get(col, "") for col in SURVEY_COLUMNS])

    # --- choices sheet ---
    ws2 = wb.create_sheet("choices")
    ws2.append(CHOICES_COLUMNS)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for row in choices_rows:
        ws2.append(list(row))

    # --- settings sheet ---
    ws3 = wb.create_sheet("settings")
    ws3.append(SETTINGS_COLUMNS)
    for cell in ws3[1]:
        cell.font = Font(bold=True)
    ws3.append([settings_row[c] for c in SETTINGS_COLUMNS])

    # cosmetic: reasonable column widths + wrap on label/note columns
    for sheet in (ws, ws2, ws3):
        for col_cells in sheet.columns:
            length = max((len(str(c.value)) if c.value is not None else 0)
                         for c in col_cells)
            col_letter = get_column_letter(col_cells[0].column)
            sheet.column_dimensions[col_letter].width = min(max(length + 2, 12), 60)
    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")  # label col

    out_path = "criminal_risk_assessment.xlsx"
    wb.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    build()
